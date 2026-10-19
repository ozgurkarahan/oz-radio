"""Content quality of the generated scripts (finding F5).

The whole source reaches the model (a fair budget per slide when a deck is
bigger than the cap), the prompts ground every claim in the source at every
audience level, never let a pre-recorded episode pass as live, keep one
language and end with a wrap-up; the deep levels get a turn count and a
coverage target that scale with the deck; sampling is conservative.
No model call: the OpenAI client is replaced by a recorder.
"""
from __future__ import annotations

import asyncio
import json
import logging
from pathlib import Path
from types import SimpleNamespace

import pytest

from services import podcast_script as ps
from services import video_script as vs
from services.podcast_models import Document, ScriptRequest, Section
from services.video_models import ScriptRequest as VideoScriptRequest, SlideRef, VideoDocument

LEVELS = ["beginner", "practitioner", "expert"]
VOICE = "en-US-Ava:DragonHDLatestNeural"


def _find_expert_deck() -> Path | None:
    """The committed Expert deck: ``content/ai-enablement`` in the repo, ``maintenance/decks`` in the image."""
    name = "ai-enablement-3-expert.pptx"
    for parent in Path(__file__).resolve().parents:
        for candidate in (parent / "content" / "ai-enablement" / name, parent / "maintenance" / "decks" / name):
            if candidate.is_file():
                return candidate
    return None


EXPERT_DECK = _find_expert_deck()


# ---------------------------------------------------------------------------
# Fixtures: decks with a unique closing sentence per slide, so that a cut is visible
# ---------------------------------------------------------------------------
def _prose(tag: str, chars: int) -> str:
    """About ``chars`` characters of sentences that end with a sentence unique to ``tag``."""
    end = f"The end marker is {tag}."
    out, k = "", 0
    while len(out) + len(end) + 1 < chars:
        k += 1
        out += f"Sentence {k} of {tag} explains one idea of the topic in plain words. "
    return out + end


def _deck(n: int, notes_chars: int = 930, body_chars: int = 350, notes_of: dict[int, int] | None = None) -> Document:
    """A pptx-shaped Document: the section text is "body\\nnotes", as ``ingest_pptx`` stores it."""
    titles = [f"Topic {i}" for i in range(n)]
    bodies = ["\n".join(f"Bullet {j} of B{i}" for j in range(4)) + "\n" + _prose(f"B{i}", body_chars)
              for i in range(n)]
    notes = [_prose(f"N{i}", (notes_of or {}).get(i, notes_chars)) for i in range(n)]
    return Document(
        id="deck", title="Deck", source_kind="pptx", slide_titles=titles, slide_notes=notes,
        sections=[Section(heading=t, text=f"{b}\n{nt}") for t, b, nt in zip(titles, bodies, notes)],
    )


def _legacy_excerpt(doc: Document) -> str:
    """The excerpt as built before the fix (notes cut at 400, then the whole text cut at 8,000)."""
    parts = ["## Slide outline"]
    for i, (title, notes) in enumerate(zip(doc.slide_titles, doc.slide_notes)):
        parts.append(f"[slide {i}] {title}")
        parts.append(f"    notes: {notes.strip()[:400]}")
    for s in doc.sections:
        parts.append(f"\n## {s.heading}")
        parts.append(s.text)
    return "\n".join(parts)


def _slides_doc(n: int) -> Document:
    return Document(id="d", title="T", source_kind="pptx", slide_titles=[f"S{i}" for i in range(n)],
                    slide_notes=[""] * n, sections=[Section(heading=f"S{i}", text="x") for i in range(n)])


def _video_deck(n: int, preview_chars: int) -> VideoDocument:
    return VideoDocument(doc_id="v", title="Deck", source_kind="pptx", slides=[
        SlideRef(index=i, image_ref=f"s{i}", title=f"Topic {i}", preview_text=_prose(f"V{i}", preview_chars))
        for i in range(n)])


# ---------------------------------------------------------------------------
# The source budget
# ---------------------------------------------------------------------------
def test_a_21k_char_deck_reaches_the_model_whole():
    doc = _deck(12)
    legacy = _legacy_excerpt(doc)
    assert 20_000 < len(legacy) < 23_000     # the Expert deck was 21,329 characters
    seen_before = legacy[:8000]
    assert "The end marker is N5." not in seen_before and "The end marker is B11." not in seen_before

    excerpt = ps._build_source_excerpt(doc)
    assert len(excerpt) <= ps.SOURCE_MAX_CHARS == 48_000
    for i in range(12):
        assert f"[slide {i}] Topic {i}" in excerpt
        assert f"The end marker is N{i}." in excerpt        # last sentence of the notes (was cut at 400)
        assert f"The end marker is B{i}." in excerpt        # last sentence of the slide text
    assert "…" not in excerpt and ps._SHORTENED_NOTE not in excerpt
    assert all(kept == source for _, source, kept in ps.source_budget(doc))


@pytest.mark.skipif(EXPERT_DECK is None, reason="the Expert deck is not present")
def test_the_real_expert_deck_reaches_the_model_whole():
    import io

    from pptx import Presentation

    from services.pptx_parser import _extract_body, _extract_notes, _extract_title

    prs = Presentation(io.BytesIO(EXPERT_DECK.read_bytes()))
    sections, titles, notes = [], [], []
    for i, slide in enumerate(prs.slides):
        title, body, note = _extract_title(slide), _extract_body(slide), _extract_notes(slide)
        sections.append(Section(heading=title or f"Slide {i + 1}", text="\n".join(p for p in (body, note) if p)))
        titles.append(title)
        notes.append(note)
    doc = Document(id="expert", title="Under the hood · Expert", source_kind="pptx", sections=sections,
                   slide_titles=titles, slide_notes=notes)
    assert len(titles) == 12 and len(_legacy_excerpt(doc)) > 20_000

    excerpt = ps._build_source_excerpt(doc)
    assert len(excerpt) <= ps.SOURCE_MAX_CHARS
    for i, note in enumerate(notes):
        assert " ".join(note.split())[-80:] in excerpt, f"notes of slide {i} were cut"
    assert all(kept == source for _, source, kept in ps.source_budget(doc))

    request = ScriptRequest(document_id="expert", num_turns=8, audience_level="expert")   # as maintenance passes it
    prompt = ps._build_system_prompt(request, doc)
    assert ps.effective_num_turns(request, doc) == 21 and "about 21 turns" in prompt
    assert "Cover at least 10 of the 12 slides" in prompt


def test_a_deck_bigger_than_the_cap_gives_every_slide_a_fair_share():
    doc = _deck(60, notes_chars=1250, body_chars=100)           # about 85,000 characters of source
    excerpt = ps._build_source_excerpt(doc)
    assert len(excerpt) <= ps.SOURCE_MAX_CHARS and ps._SHORTENED_NOTE in excerpt
    for i in range(60):
        assert f"[slide {i}] Topic {i}" in excerpt
        assert f"The end marker is B{i}." in excerpt         # a short slide text stays whole
        assert f"Sentence 1 of N{i} " in excerpt             # every slide's notes start; the tail is not dropped
    kept = [k for _, _, k in ps.source_budget(doc)]
    assert min(kept) >= 600 and max(kept) - min(kept) <= 120
    assert sum(kept) <= ps.SOURCE_MAX_CHARS
    assert "Sentence 1 of N59 " not in _legacy_excerpt(doc)[:8000]      # the old cut never reached it


def test_one_huge_slide_does_not_starve_the_others():
    doc = _deck(12, notes_of={3: 40_000})
    excerpt = ps._build_source_excerpt(doc)
    assert len(excerpt) <= ps.SOURCE_MAX_CHARS
    for i in (0, 1, 2, 4, 5, 6, 7, 8, 9, 10, 11):
        assert f"The end marker is N{i}." in excerpt         # the short slides stay whole
    assert "The end marker is N3." not in excerpt and "[slide 3] Topic 3" in excerpt
    budget = ps.source_budget(doc)
    assert budget[3][2] > 30_000 and all(kept == source for i, (_, source, kept) in enumerate(budget) if i != 3)
    assert excerpt.count("…") >= 1


def test_a_source_without_slides_shares_the_budget_between_its_sections():
    doc = Document(id="r", title="Report", source_kind="pdf",
                   sections=[Section(heading=f"Page {i}", text=_prose(f"P{i}", 40_000)) for i in range(3)])
    excerpt = ps._build_source_excerpt(doc)
    assert len(excerpt) <= ps.SOURCE_MAX_CHARS and "## Slides" not in excerpt
    for i in range(3):
        assert f"## Page {i}" in excerpt and f"Sentence 1 of P{i} " in excerpt
    assert ps.source_budget(doc)[2][2] > 14_000


def test_misaligned_slide_lists_do_not_break_the_excerpt():
    doc = Document(id="d", title="T", source_kind="pptx", slide_titles=["Intro", "Next", "Last"],
                   slide_notes=["n"], sections=[Section(text="x")])
    excerpt = ps._build_source_excerpt(doc)
    assert "[slide 0] Intro" in excerpt and "on slide: x" in excerpt and "notes: n" in excerpt
    assert "[slide 1] Next" in excerpt and "[slide 2] Last" in excerpt


def test_clip_text_cuts_at_a_boundary_and_marks_the_cut():
    text = "One two three. Four five six seven. Eight nine ten eleven twelve."
    assert ps.clip_text(text, 200) == text
    assert ps.clip_text(text, 40) == "One two three. Four five six seven.…"
    assert ps.clip_text("Alpha beta gamma delta epsilon zeta", 20) == "Alpha beta gamma…"
    assert ps.clip_text("abcdefghijkl", 5) == "abcd…" and ps.clip_text("abc", 1) == "…" and ps.clip_text("abc", 0) == ""
    assert ps.clip_text("第一句话。第二句话。第三句话。", 12) == "第一句话。第二句话。…"     # no spaces in CJK text
    for limit in range(2, len(text)):
        assert len(ps.clip_text(text, limit)) <= limit


@pytest.mark.parametrize("sizes,budget,cap", [
    ([], 10, 0), ([10, 20], 100, 20), ([10, 20, 30, 1000], 100, 40), ([500, 500], 600, 300),
    ([5, 5, 5], 0, 0), ([1000, 1000, 10], 210, 100),
])
def test_fair_share_cap(sizes, budget, cap):
    assert ps.fair_share_cap(sizes, budget) == cap
    assert sum(min(s, cap) for s in sizes) <= max(budget, 0) or sum(sizes) <= budget


# ---------------------------------------------------------------------------
# Turns and coverage
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("level,slides,asked,expected", [
    ("expert", 12, 8, 21),          # 10 exchanges (a question and an answer) + a wrap-up turn
    ("expert", 12, 16, 21),         # the caller's number stays a floor, never a ceiling
    ("practitioner", 12, 8, 17),
    ("beginner", 12, 8, 8),         # short and simple: never scaled
    (None, 12, 8, 8),
    ("expert", 2, 8, 8),
    ("expert", 5, 8, 9),
    ("practitioner", 5, 16, 16),
    ("expert", 100, 8, ps.MAX_TARGET_TURNS),
    ("beginner", 100, 8, 8),
])
def test_effective_num_turns_scales_with_the_slide_count(level, slides, asked, expected):
    req = ScriptRequest(document_id="d", num_turns=asked, audience_level=level)
    assert ps.effective_num_turns(req, _slides_doc(slides)) == expected


def test_the_turn_target_leaves_room_below_the_render_limit():
    from security import MAX_TURNS_PER_SCRIPT

    assert ps.MAX_TARGET_TURNS * 5 // 4 < MAX_TURNS_PER_SCRIPT     # a model overshoot of a quarter still renders


@pytest.mark.parametrize("level,slides,turns,target", [
    ("expert", 12, 21, 10), ("practitioner", 12, 17, 8), ("expert", 100, 30, 14),
    ("beginner", 12, 8, 0), (None, 12, 8, 0), ("expert", 0, 8, 0),
])
def test_coverage_target(level, slides, turns, target):
    assert ps.coverage_target(level, slides, turns) == target


def test_coverage_text_per_level_for_a_12_slide_deck():
    doc = _slides_doc(12)

    def prompt(level):
        return ps._build_system_prompt(ScriptRequest(document_id="d", num_turns=8, audience_level=level), doc)

    expert, practitioner, beginner, generic = (prompt("expert"), prompt("practitioner"), prompt("beginner"),
                                               prompt(None))
    assert "Cover at least 10 of the 12 slides" in expert and "about 21 turns" in expert
    assert "full deep dive" in expert and "one precise question, one precise answer" in expert
    assert "cover at least 8 of the 12 slides" in practitioner and "about 17 turns" in practitioner
    # No bloat at beginner level: a few ideas, the requested length.
    assert "keep the episode short and simple" in beginner and "3 to 5 most important ideas" in beginner
    assert "about 8 turns" in beginner and "of the 12 slides" not in beginner
    assert "cover the main points of the slides in order" in generic and "about 8 turns" in generic
    assert "of the 12 slides" not in generic


def test_coverage_text_without_slides():
    doc = Document(id="r", title="R", source_kind="pdf", sections=[Section(heading="Page 1", text="x")])

    def prompt(level):
        return ps._build_system_prompt(ScriptRequest(document_id="d", num_turns=8, audience_level=level), doc)

    assert "Cover every main section of the source" in prompt("expert") and "about 8 turns" in prompt("expert")
    assert "cover the main sections of the source in order" in prompt("practitioner")
    assert "short and simple" in prompt("beginner") and "cover the main points of the source" in prompt(None)


# ---------------------------------------------------------------------------
# Prompt rules (every level, every language)
# ---------------------------------------------------------------------------
GROUNDING = ["GROUNDING", "Use ONLY what the source below states", "do not fill gaps", "Do NOT invent statistics",
             "mitigations", "anecdotes", "hedges", "\"at some cost to accuracy\" must never become \"almost no loss\"",
             "a limit must not turn into a guarantee", "never guess what was cut"]
RECORDED = ["RECORDED EPISODE", "pre-recorded", "Never say or imply that it is live", "unscripted", "generated on the fly",
            "\"nothing is recorded in advance\"", "Do not invite listeners to ask the hosts questions"]
WRAP_UP = ["WRAP-UP", "The final turn must be a short wrap-up", "recaps the key points", "no new facts"]


@pytest.mark.parametrize("language,name", [("en-US", "English"), ("fr-FR", "French"), ("es-ES", "Spanish")])
@pytest.mark.parametrize("level", LEVELS + [None])
def test_every_level_carries_the_grounding_recorded_language_and_wrap_up_rules(level, language, name):
    prompt = ps._build_system_prompt(
        ScriptRequest(document_id="d", language=language, audience_level=level), _slides_doc(3))
    for marker in GROUNDING + RECORDED + WRAP_UP:
        assert marker in prompt, marker
    assert f"must be entirely in {name}" in prompt and "never slip into another language" in prompt
    # The rules come before the output format, which stays the last thing the model reads.
    assert prompt.index("GROUNDING") < prompt.index("RECORDED EPISODE") < prompt.index("Output format")
    assert prompt.endswith("Return ONLY the JSON object, no prose before or after.")


def test_the_debate_style_cannot_invent_evidence():
    assert "evidence from the source" in ps.STYLE_DIRECTIVES["debate"]
    assert "limits and trade-offs the source itself states" in ps.STYLE_DIRECTIVES["debate"]


def test_level_directive_does_not_loosen_the_rules():
    assert "grounding, language and output-format rules below still apply" in ps.AUDIENCE_SCOPE_NOTE


@pytest.mark.parametrize("level", LEVELS + [None])
def test_video_prompt_carries_the_grounding_recorded_and_language_rules(level):
    prompt = vs._build_system_prompt(
        VideoScriptRequest(voice=VOICE, language="fr-FR", audience_level=level), _video_deck(3, 100))
    for marker in ("GROUNDING", "Use ONLY what the slides and their notes state", "Do NOT invent statistics",
                   "hedges", "\"at some cost to accuracy\" must never become \"almost no loss\"",
                   "\"after deployment\" must not become \"before deployment\"",
                   "RECORDED VIDEO", "pre-recorded", "never say or imply that the narration is live".capitalize(),
                   "generated on the fly", "Do not invite viewers to ask questions",
                   "must be entirely in French", "never slip into another language"):
        assert marker in prompt, marker
    assert "EXACTLY 3 objects" in prompt and prompt.endswith("Return ONLY the JSON object, no prose before or after.")
    # Nothing invites the model to add content of its own any more.
    assert "add the why" not in prompt
    assert "concrete example taken from the slide or its notes" in vs.STYLE_DIRECTIVES["explainer"]


# ---------------------------------------------------------------------------
# The request sent to the model
# ---------------------------------------------------------------------------
def _chunk(text: str):
    return SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content=text))])


def _recorder(captured: dict, reply: str):
    class Completions:
        def create(self, **kwargs):
            captured.update(kwargs)
            return iter([_chunk(reply[i:i + 50]) for i in range(0, len(reply), 50)])

    return SimpleNamespace(chat=SimpleNamespace(completions=Completions()))


def _collect(agen) -> list:
    async def run():
        return [item async for item in agen]

    return asyncio.run(run())


CFG = SimpleNamespace(openai_chat_deployment="")


def _turns_reply(slide_idxs: list[int]) -> str:
    return json.dumps({"turns": [
        {"idx": i, "speaker": "interviewer" if i % 2 == 0 else "expert", "text": f"Turn {i}. Fine.", "slide_idx": s}
        for i, s in enumerate(slide_idxs)]})


def test_podcast_request_uses_the_whole_source_low_temperature_and_room_for_a_long_episode(monkeypatch):
    captured: dict = {}
    monkeypatch.setattr(ps, "get_openai_client", lambda cfg: _recorder(captured, _turns_reply([0, 0, 1])))
    doc = _deck(12)
    req = ScriptRequest(document_id="d", num_turns=8, style="debate", audience_level="expert")
    turns = _collect(ps.generate_script_stream(CFG, doc, req))
    assert [t.idx for t in turns] == [0, 1, 2]
    assert captured["temperature"] == ps.TEMPERATURE == 0.4
    assert captured["max_completion_tokens"] == ps.MAX_COMPLETION_TOKENS == 8192
    assert captured["response_format"] == {"type": "json_object"} and captured["stream"] is True
    assert captured["model"] == "gpt-4.1"
    system, user = (m["content"] for m in captured["messages"])
    assert system == ps._build_system_prompt(req, doc) and "about 21 turns" in system
    for i in range(12):
        assert f"The end marker is N{i}." in user and f"The end marker is B{i}." in user
    assert user.startswith("Here is the source material.") and "Document title: Deck" in user


def test_podcast_warns_when_a_deep_episode_covers_too_few_slides(monkeypatch, caplog):
    monkeypatch.setattr(ps, "get_openai_client", lambda cfg: _recorder({}, _turns_reply([0, 0, 1, 1, 2])))
    req = ScriptRequest(document_id="d", num_turns=8, audience_level="expert")
    with caplog.at_level(logging.WARNING, logger=ps.log.name):
        _collect(ps.generate_script_stream(CFG, _deck(12), req))
    assert any("3 of 12 slides covered, target 10" in r.getMessage() for r in caplog.records)


def test_podcast_is_quiet_when_the_coverage_target_is_met(monkeypatch, caplog):
    monkeypatch.setattr(ps, "get_openai_client",
                        lambda cfg: _recorder({}, _turns_reply([0, 0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11])))
    req = ScriptRequest(document_id="d", num_turns=8, audience_level="expert")
    with caplog.at_level(logging.WARNING, logger=ps.log.name):
        _collect(ps.generate_script_stream(CFG, _deck(12), req))
    assert not [r for r in caplog.records if "slides covered" in r.getMessage()]


def test_podcast_has_no_coverage_target_below_the_deep_levels(monkeypatch, caplog):
    monkeypatch.setattr(ps, "get_openai_client", lambda cfg: _recorder({}, _turns_reply([0, 0, 1])))
    req = ScriptRequest(document_id="d", num_turns=8, audience_level="beginner")
    with caplog.at_level(logging.WARNING, logger=ps.log.name):
        _collect(ps.generate_script_stream(CFG, _deck(12), req))
    assert not [r for r in caplog.records if "slides covered" in r.getMessage()]


def test_video_outline_keeps_the_whole_preview_of_each_slide():
    outline = vs._build_slide_outline(_video_deck(12, 600))
    for i in range(12):
        assert f"[slide {i}] Topic {i}" in outline and f"The end marker is V{i}." in outline   # [:500] cut it
    assert "…" not in outline


def test_video_outline_shares_the_budget_between_the_slides_of_a_big_deck():
    doc = _video_deck(20, 5000)                                        # 100,000 characters of previews
    outline = vs._build_slide_outline(doc)
    assert len(outline) <= vs.SOURCE_MAX_CHARS == ps.SOURCE_MAX_CHARS
    for i in range(20):
        assert f"[slide {i}] Topic {i}" in outline and f"Sentence 1 of V{i} " in outline
    sizes = [len(block) for block in outline.split("[slide ")[1:]]
    assert max(sizes) - min(sizes) <= 150 and min(sizes) > 2000
    assert vs._build_slide_outline(doc, 8000).count("[slide ") == 20      # still every slide under a small cap


def test_video_request_uses_low_temperature_and_room_for_a_long_deck(monkeypatch):
    captured: dict = {}
    reply = json.dumps({"narrations": [{"slide_index": 0, "narration": "Hello."}]})
    monkeypatch.setattr(vs, "get_openai_client", lambda cfg: _recorder(captured, reply))
    doc = _video_deck(12, 600)
    out = _collect(vs.generate_script_stream(CFG, doc, VideoScriptRequest(voice=VOICE, audience_level="expert")))
    assert [n.slide_index for n in out] == [0]
    assert captured["temperature"] == vs.TEMPERATURE == 0.4
    assert captured["max_completion_tokens"] == vs.MAX_COMPLETION_TOKENS == 8192
    assert captured["response_format"] == {"type": "json_object"} and captured["stream"] is True
    user = captured["messages"][1]["content"]
    assert "Slide count: 12" in user and all(f"The end marker is V{i}." in user for i in range(12))
